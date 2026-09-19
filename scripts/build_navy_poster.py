"""Fill the supplied SciFig template, preserving its layout and editable text.
Run from the repository root; needs python-pptx and Poppler (pdftoppm).
"""
from pathlib import Path
import subprocess
import tempfile
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
prs = Presentation(ROOT / 'scifig-a0-portrait-poster-template-classic-navy.pptx')
slide = prs.slides[1]
shapes = {s.name: s for s in slide.shapes}
NAVY = '17365D'
INK = '253446'


def text(name, paragraphs, size=30, color=INK, bold_first=False):
    shape = shapes[name]
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    if isinstance(paragraphs, str):
        paragraphs = [paragraphs]
    for i, value in enumerate(paragraphs):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = value
        p.font.name = 'Arial'
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor.from_string(color)
        p.font.bold = bold_first and i == 0
        p.space_after = Pt(size * 0.7)
        p.line_spacing = 1.14
    return shape


def remove(name):
    element = shapes[name]._element
    element.getparent().remove(element)


def figure(name, filename):
    box = shapes[name]
    if name in ('Picture Placeholder 5', 'Picture Placeholder 6', 'Picture Placeholder 7'):
        box.width = Inches(11.01)
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / 'figure'
        subprocess.run(['pdftoppm', '-r', '240', '-png', '-singlefile',
                        str(ROOT / 'figures' / filename), str(out)], check=True)
        img = Image.open(str(out) + '.png')
        ratio = min(box.width / img.width, box.height / img.height)
        width, height = round(img.width * ratio), round(img.height * ratio)
        picture = slide.shapes.add_picture(str(out) + '.png',
            box.left + (box.width-width)//2,
            box.top + (box.height-height)//2, width, height)
        picture.name = 'Project figure: ' + filename
    remove(name)

# Only the finished poster; the template itself remains untouched.
first = prs.slides._sldIdLst[0]
prs.part.drop_rel(first.rId)
prs.slides._sldIdLst.remove(first)

text('poster-title', ['Non-Smooth Equilibria in Hill-Type Muscle Models'],
     size=76, color='FFFFFF', bold_first=True)
text('poster-authors', 'Alan Royce Gabriel Samuel', size=48, color='FFFFFF')
text('poster-affiliations', 'Indian Institute of Technology Madras, Chennai, India\nIROS 2026 · Workshop on Neuromuscular Robotics', size=30, color='FFFFFF')
remove('Picture Placeholder 57')
remove('logo-backing')
text('logo-hint', 'IIT MADRAS\nPOSTER 5', size=36, color='FFFFFF', bold_first=True)
shapes['logo-hint'].top = Inches(1.45)
shapes['logo-hint'].height = Inches(1.5)
for p in shapes['logo-hint'].text_frame.paragraphs:
    p.alignment = PP_ALIGN.CENTER

text('Text Placeholder 3', [
    'One equilibrium. Four directional linear models.',
    'Muscle-driven robots, prostheses and exosuits often use controllers designed around an equilibrium linearization. This assumes that a single local model describes small perturbations.',
    'The Hill-type formulation studied here switches twice: activation dynamics changes when excitation crosses activation, and the force–velocity relation changes between shortening and lengthening.',
    'At equilibrium, excitation equals activation and velocity is zero. Both switching surfaces therefore pass through the operating point.',
    'Question: how much does the choice of branch affect prediction and closed-loop control?'
], size=30, bold_first=True)

text('Text Placeholder 11', [
    'A simulated elbow against gravity',
    'A forearm–hand pendulum is driven by a rigid-tendon Hill-type brachialis model. Muscle parameters follow Holzbaur; activation dynamics follows Thelen; moment arms are digitized from Murray et al.',
    'Reference posture: 60°\nEquilibrium activation: 0.113',
    'Prediction test: apply a small +0.01 excitation step and compare all four local models with the nonlinear trajectory.',
    'Feedback test: linear MPC with a 10 ms sample time and 20-step horizon. Keep controller weights fixed; vary only the branch and linearization point.',
    'Extensions: held-load equilibria, a brachialis–triceps antagonist pair, noise, plant mismatch and delay.'
], size=30, bold_first=True)

text('Text Placeholder 12', [
    '14–75% prediction error from a wrong branch',
    'At the reference posture, matched-branch error stays near 1% of the true angular swing. Mismatched models diverge within 50–100 ms.',
    'At 50 ms: wrong activation branch, 71%; wrong velocity branch, 14%; both wrong, 75%. These are errors relative to the true swing, not absolute angles.',
    'The effect persists across operating points',
    'With both branches wrong, error is at least 23% of the activating-step signal at every single-muscle equilibrium tested.',
    'Feedback changes the consequence',
    'In the tested delay-free cases, feedback preserves stability and steady-state tracking, but differences in transient quality remain. Blending the branches gives 8–9× more raising overshoot than the fixed, sign-test and relinearized designs.'
], size=32, bold_first=True)

text('Text Placeholder 13', 'Read together: branch choice affects prediction across operating points and remains visible in closed-loop transients.', size=26)

text('Text Placeholder 14', [
    'Relinearize at the measured state',
    'During lowering, braking re-activates the muscle. A branch chosen only from the initial sign misses this mid-plan change.',
    'Lowering overshoot falls from 0.37° with the sign test to 0.02° with relinearization.',
    'Smoothing activation retains much of the transient cost. Tracking rankings depend on weights and noise.',
    'A predictor restored stability under the tested 20 and 50 ms delays.'
], size=29, bold_first=True)
text('results-table-caption', 'Baseline MPC: overshoot and tracking error (degrees).', size=21)
table = shapes['results-table'].table
for col, width in zip(table.columns, [2.8, 1.6, 1.6, 2.21]):
    col.width = Inches(width)
rows = [
    ['Model', 'Up', 'Down', 'Sine RMS'],
    ['Blended', '0.65', '0.15', '0.196'],
    ['Smooth act.', '0.42', '0.48', '0.169'],
    ['Sign test', '0.08', '0.37', '0.153'],
    ['Relinearize', '0.07', '0.02', '0.149'],
]
for i, row in enumerate(rows):
    for j, value in enumerate(row):
        cell = table.cell(i,j)
        cell.text = value
        for p in cell.text_frame.paragraphs:
            p.font.name = 'Arial'
            p.font.size = Pt(23)
            p.font.bold = i in (0,4)
            p.font.color.rgb = RGBColor.from_string('FFFFFF' if i == 0 else INK)
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER

a = text('Text Placeholder 15', [
    'Choose the branch as well as the operating point.',
    'Use directional local models and relinearize as the state changes. Co-contraction can reduce the velocity kink, but adds an activation switch and effort.',
    'Scope: one simulated joint. Hardware validation, reflexes and history-dependent muscle effects remain open.'
], size=28, bold_first=True)
text('Text Placeholder 16', [
    'Thelen (2003). J Biomech Eng 125:70–77.',
    'Katz (1939). J Physiol 96:45–64.',
    'Holzbaur et al. (2005). Ann Biomed Eng 33:829–840.',
    'Murray et al. (1995). J Biomech 28:513–525.',
    'Camlibel et al. (2008). Automatica 44:1261–1267.',
    'De Groote et al. (2016). Ann Biomed Eng 44:2922–2936.'
], size=23)
text('sec-acknow-title', 'CONTACT & MATERIALS', size=40, color='FFFFFF', bold_first=True)
text('Text Placeholder 17', [
    'Alan Royce Gabriel Samuel',
    'bs22b001@smail.iitm.ac.in',
    'Poster source and figures:\ngithub.com/alanroyce2010/hill_IROS'
], size=26, bold_first=True)
remove('qr-slot')
text('qr-slot-hint', 'Simulation results from the existing project draft.', size=22)
text('footer-credit', 'IROS 2026 · 1st Workshop on Neuromuscular Robotics · Poster 5   |   Layout: SciFig AI Classic Navy', size=24, color=NAVY)

for slot, filename, caption_name, caption in [
    ('Picture Placeholder 2', 'fig1_kinks.pdf', 'fig-1-caption', 'Activation slopes differ by 7.45× at the reference equilibrium; velocity slopes differ by 2×.'),
    ('Picture Placeholder 4', 'fig6_validation.pdf', 'fig-2-caption', 'Model checks: moment arm and torque–angle relationship. Brachialis peaks at 25.8 N m near 94°.'),
    ('Picture Placeholder 5', 'fig2_prediction.pdf', 'fig-3-caption', 'A small excitation step: the matched branch tracks the nonlinear model; wrong branches diverge.'),
    ('Picture Placeholder 6', 'fig3_equilibria.pdf', 'fig-4-caption', 'Held-load sweep: activation-branch error falls as activation rises; velocity-branch error grows.'),
    ('Picture Placeholder 7', 'fig4_closedloop.pdf', 'fig-5-caption', 'Closed-loop steps: 60° → 63° → 57°. Relinearization reduces lowering-step overshoot.'),
    ('Picture Placeholder 8', 'fig5_pair.pdf', 'fig-6-caption', 'Antagonist co-contraction reduces the velocity kink toward cancellation.'),
]:
    figure(slot, filename)
    text(caption_name, caption, size=22)
    if caption_name in ('fig-3-caption', 'fig-4-caption', 'fig-5-caption'):
        shapes[caption_name].width = Inches(11.01)

slide.notes_slide.notes_text_frame.text = ('Visual draft populated from the existing muscle-activation-control poster. '
    'Numerical results have not been independently revalidated. Original SciFig template preserved separately. '
    'Images fitted without cropping. Text and results table remain editable.')
output = ROOT / 'muscle-control-classic-navy.pptx'
prs.save(output)
print(output)
