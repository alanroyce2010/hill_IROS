"""Build an editable A0 navy poster in the requested section order.
Abstract condensed from the ICRA 2027 submission; original retained in content/.
Requires python-pptx, Pillow and Poppler. Run from any directory.
"""
from pathlib import Path
from copy import deepcopy
import subprocess
import tempfile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

ROOT=Path(__file__).resolve().parents[1]
prs=Presentation(ROOT/'scifig-a0-portrait-poster-template-classic-navy.pptx')
slide=prs.slides[1]
original={s.name:s for s in slide.shapes}
NAVY='1D3A60'; BLUE='386CB1'; INK='253446'; LIGHT='EDF3FA'
# Use the actual template's panel and heading styles.
panel_xml=deepcopy(original['sec-introd']._element)
bar_xml=deepcopy(original['sec-introd-head']._element)
for s in list(slide.shapes):
    if s.name!='header-band':
        s._element.getparent().remove(s._element)
first=prs.slides._sldIdLst[0]
prs.part.drop_rel(first.rId);prs.slides._sldIdLst.remove(first)


def textbox(name,x,y,w,h,paras,size=28,color=INK,bold=False,align=PP_ALIGN.LEFT,gap=12):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));shape.name=name
    f=shape.text_frame;f.clear();f.word_wrap=True
    f.margin_left=f.margin_right=Inches(.03);f.margin_top=f.margin_bottom=0
    f.vertical_anchor=MSO_ANCHOR.TOP
    if isinstance(paras,str):paras=[paras]
    for i,txt in enumerate(paras):
        p=f.paragraphs[0] if i==0 else f.add_paragraph()
        p.text=txt;p.alignment=align;p.font.name='Arial';p.font.size=Pt(size)
        p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(gap);p.line_spacing=1.1
    return shape


def clone(xml,name,x,y,w,h):
    el=deepcopy(xml);slide.shapes._spTree.insert_element_before(el,'p:extLst')
    s=slide.shapes[-1];s.shape_id # force wrapper
    nv=el.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
    nv.set('id',str(max(t.shape_id for t in slide.shapes)+1));nv.set('name',name)
    s.left=Inches(x);s.top=Inches(y);s.width=Inches(w);s.height=Inches(h)
    return s


def panel(title,x,y,w,h):
    clone(panel_xml,title+' panel',x,y,w,h)
    clone(bar_xml,title+' bar',x,y,w,1.02)
    textbox(title+' heading',x+.2,y+.18,w-.4,.75,title,39,'FFFFFF',True,gap=0)


def figure(file,x,y,w,h):
    with tempfile.TemporaryDirectory() as tmp:
        target=Path(tmp)/'figure'
        subprocess.run(['pdftoppm','-r','240','-png','-singlefile',str(ROOT/'figures'/file),str(target)],check=True)
        im=Image.open(str(target)+'.png');ratio=min(w/im.width,h/im.height)
        iw,ih=im.width*ratio,im.height*ratio
        s=slide.shapes.add_picture(str(target)+'.png', Inches(x+(w-iw)/2), Inches(y+(h-ih)/2),Inches(iw),Inches(ih));s.name=file


def box(x,y,w,h,title,body):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(LIGHT)
    s.line.color.rgb=RGBColor.from_string(BLUE)
    textbox(title,x+.12,y+.2,w-.24,.7,title,27,NAVY,True,PP_ALIGN.CENTER,gap=0)
    textbox(title+' detail',x+.12,y+1.02,w-.24,h-1.1,body,25,INK,False,PP_ALIGN.CENTER,gap=0)

# Template header and full scientific title.
textbox('title',1.1,.55,30.9,1.65,'Non-Smooth Equilibria in Hill-Type Muscle Models',70,'FFFFFF',True,PP_ALIGN.CENTER)
textbox('subtitle',1.1,1.95,30.9,.85,'A Quantified Consequence for Linearization-Based Control',43,'FFFFFF',False,PP_ALIGN.CENTER)
textbox('author',1.1,3.02,30.9,.65,'Alan Royce Gabriel Samuel · Indian Institute of Technology Madras',36,'FFFFFF',False,PP_ALIGN.CENTER)
textbox('event',1.1,3.83,30.9,.6,'IROS 2026 · Workshop on Neuromuscular Robotics · Poster 5',28,'FFFFFF',False,PP_ALIGN.CENTER)

abstract=(ROOT/'content/poster-abstract.txt').read_text().strip()
panel('ABSTRACT',1.16,5.3,30.79,3.6)
textbox('concise abstract',1.4,6.55,30.31,2.1,abstract,32,gap=0)

panel('INTRODUCTION',1.16,11.85,10.0,6.95)
textbox('introduction',1.4,13.14,9.52,5.4,[
    'Muscle models connect excitation to force in prostheses, exoskeletons, FES and tendon-driven robots. LQR and successive-linearization MPC use a local Jacobian to predict motion [1,2].',
    'But an equilibrium lies on two switches: excitation equals activation, and fibre velocity is zero. The perturbation direction selects one of four local models.',
    'We quantify the prediction error and test how much of it survives feedback.'
],28,gap=17)

panel('FIGURE 1 · PROJECT OVERVIEW',11.56,11.85,20.39,6.95)
# Editable overview: physiological model, directional linearization, and controller.
for x,t,b in [
    (11.87,'MUSCLE + JOINT','Excitation → activation\n→ force → elbow angle'),
    (16.89,'TWO SWITCHES','u* = a*\nv* = 0'),
    (21.91,'FOUR MODELS','Activation branch ×\nvelocity branch'),
    (26.93,'CONTROL TEST','Prediction horizon\n+ closed-loop MPC'),
]:
    box(x,13.55,4.55,2.8,t,b)
for x in [16.46,21.48,26.50]:
    arrow=slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x),Inches(14.65),Inches(.4),Inches(.45))
    arrow.fill.solid();arrow.fill.fore_color.rgb=RGBColor.from_string(BLUE);arrow.line.fill.background()
textbox('overview takeaway',12.0,16.7,19.5,.7,'Wrong branch: 14–75% prediction error  |  Relinearize: 0.37° → 0.02° lowering overshoot',29,NAVY,True,PP_ALIGN.CENTER)
textbox('overview caption',12.0,17.67,19.5,.82,'Figure 1. From equilibrium switching to controller performance. Read the active branch and update the local model at the measured state.',24,gap=0)

# Read down the left column, then the centre and right columns.
x1,x2,x3=1.16,11.56,21.96
w=9.99
panel('RELATED WORK',x1,19.2,w,8.0)
textbox('related work',x1+.24,20.47,w-.48,6.45,[
    'Hill-type modelling. Thelen models activation and force generation; Millard et al. characterize computational muscle formulations [1,2].',
    'Smoothing. Kirsch et al. smooth torque–velocity dynamics; De Groote et al. smooth activation for optimization [3,4]. Here we measure the cost of branch choice at equilibrium.',
    'Non-smooth systems. Camlibel et al. give continuity conditions for bimodal linear systems [5]. Yeo et al. study a distinct force–length instability, not the activation-switch control problem [6].'
],28,gap=17)

panel('FORMULATION / MODEL',x1,27.6,w,17.25)
textbox('model',x1+.24,28.88,w-.48,4.7,[
    'Plant: one elbow, a forearm–hand pendulum, and a rigid-tendon Hill-type brachialis. Muscle parameters follow Holzbaur; moment arms follow Murray [7,8].',
    'State x = [θ, θ̇, a]ᵀ; input u is neural excitation.',
    'Iθ̈ = Fᴹᵀ r(θ) − mglc sin θ\nv = −r(θ)θ̇ / l₀ᴹ',
],28,gap=15)
textbox('activation equation',x1+.24,33.32,w-.48,2.2,[
    'ȧ = (u − a) / τa',
    'τa = τact (0.5 + 1.5a), when u > a\nτa = τdeact / (0.5 + 1.5a), when u ≤ a'
],28,NAVY,False,gap=15)
figure('fig1_kinks.pdf',x1+.25,35.65,w-.5,3.55)
textbox('kinks caption',x1+.25,39.35,w-.5,1.0,'Figure 2. At θ* = 60° and a* = 0.113, the activation slope ratio is 7.45 and the velocity slope ratio is 2.',23,gap=0)
textbox('local models',x1+.24,40.63,w-.48,3.95,[
    'Four local models: δẋ = Aᵢδx + Bᵢδu.',
    'Each switch satisfies A₁ − A₂ = ecᵀ and B₁ − B₂ = ed [5]; the two switches affect different Jacobian entries.',
    'MPC: 10 ms sampling, 20-step horizon. Compare rules using the same plant and controller weights.'
],27,gap=14)

panel('RESULTS',x2,19.2,w,25.65)
textbox('prediction result',x2+.24,20.47,w-.48,2.35,[
    'Prediction: 14–75% wrong-branch error',
    'Matched-branch error stays near 1% at the reference posture. With both branches wrong, error is ≥23% at every tested single-muscle equilibrium for activating steps.'
],27,gap=12)
figure('fig2_prediction.pdf',x2+.25,23.16,w-.5,4.4)
textbox('prediction caption',x2+.25,27.65,w-.5,.9,'Figure 3. Response to a +0.01 excitation step: branch errors separate within 50–100 ms.',23,gap=0)
textbox('feedback result',x2+.24,28.82,w-.48,1.95,[
    'Feedback: transient quality remains',
    'Central-difference blending gives 8–9× more raising overshoot than the naive, sign-test and relinearizing rules.'
],27,gap=12)
figure('fig4_closedloop.pdf',x2+.25,31.00,w-.5,3.95)
textbox('feedback caption',x2+.25,35.06,w-.5,.9,'Figure 4. Raising and lowering steps. Updating the model at the measured state reduces lowering overshoot.',23,gap=0)
shape=slide.shapes.add_table(6,4,Inches(x2+.25),Inches(36.3),Inches(w-.5),Inches(3.7))
table=shape.table
for col,wi in zip(table.columns,[3.6,1.8,1.8,2.29]):col.width=Inches(wi)
for i,row in enumerate([
    ['Rule','Up','Down','Sine RMS'],
    ['Blended','0.65','0.15','0.196'],
    ['Smoothed act.','0.42','0.48','0.169'],
    ['Naive fixed','0.07','0.00','0.200'],
    ['Sign test','0.08','0.37','0.153'],
    ['Relinearize','0.07','0.02','0.149'],
]):
    for j,val in enumerate(row):
        c=table.cell(i,j);c.text=val;c.fill.solid();c.fill.fore_color.rgb=RGBColor.from_string(NAVY if i==0 else ('EDF3FA' if i==5 else 'FFFFFF'))
        for p in c.text_frame.paragraphs:
            p.font.name='Arial';p.font.size=Pt(23);p.font.bold=i in (0,5)
            p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else INK)
textbox('table caption',x2+.25,40.18,w-.5,.7,'Table 1. Overshoot and sine-tracking RMS, degrees.',23,gap=0)
textbox('extensions results',x2+.24,41.18,w-.48,3.3,[
    'Co-contraction can cancel the joint velocity kink, while adding a second activation switch.',
    'Unmodelled delay can cause divergence; a predictor restored stability in the tested cases. Tracking rankings depend on weights and noise.'
],27,gap=14)

panel('CONCLUSION',x3,19.2,w,10.35)
textbox('conclusion',x3+.24,20.47,w-.48,8.6,[
    'Choose the branch as well as the operating point.',
    '1. Equilibria of the studied switched Hill-type model require directional local linearizations.',
    '2. Branch mismatch causes substantial prediction errors. Feedback reduces the effect to transient quality in the tested delay-free cases.',
    '3. Relinearize at the measured state, avoid central differences across a kink, and compensate delay.',
    'Limits: simulation only; one joint, one or two muscles; no reflexes, short-range stiffness or history dependence. Hardware validation remains open.'
],28,gap=19)

panel('REFERENCES',x3,29.95,w,14.9)
refs=[
    '[1] Thelen DG (2003). Adjustment of muscle mechanics model parameters to simulate dynamic contractions in older adults. J Biomech Eng 125:70–77.',
    '[2] Millard M et al. (2013). Flexing computational muscle: modeling and simulation of musculotendon dynamics. J Biomech Eng 135:021005.',
    '[3] Kirsch NA et al. (2018). Model-based dynamic control allocation in a hybrid neuroprosthesis. IEEE TNSRE 26:224–232.',
    '[4] De Groote F et al. (2016). Evaluation of direct collocation optimal control problem formulations for solving the muscle redundancy problem. Ann Biomed Eng 44:2922–2936.',
    '[5] Camlibel MK, Heemels WPMH, Schumacher JM (2008). A full characterization of stabilizability of bimodal piecewise linear systems with scalar inputs. Automatica 44:1261–1267.',
    '[6] Yeo SH et al. (2023). Numerical instability of Hill-type muscle models. J R Soc Interface 20:20220430.',
    '[7] Holzbaur KRS, Murray WM, Delp SL (2005). A model of the upper extremity for simulating musculoskeletal surgery and analyzing neuromuscular control. Ann Biomed Eng 33:829–840.',
    '[8] Murray WM, Delp SL, Buchanan TS (1995). Variation of muscle moment arms with elbow and forearm position. J Biomech 28:513–525.'
]
textbox('references',x3+.24,31.25,w-.48,12.95,refs,24,gap=18)
textbox('footer',1.16,45.65,30.79,.7,'bs22b001@smail.iitm.ac.in  |  github.com/alanroyce2010/hill_IROS  |  Layout adapted from SciFig AI Classic Navy',24,NAVY,False,PP_ALIGN.CENTER,gap=0)
# Close the space freed by the shorter abstract while retaining the bottom margin.
for shape in slide.shapes:
    if shape.top >= Inches(11.8) and shape.name != 'footer':
        shape.top -= Inches(2.55)
    if shape.name in ('FORMULATION / MODEL panel', 'RESULTS panel', 'REFERENCES panel'):
        shape.height += Inches(2.55)

slide.notes_slide.notes_text_frame.text='Abstract condensed from ICRA 2027 main.tex. Figure 1 is an editable conceptual overview. Numerical results are from the existing submission; no new experiments were run.'
prs.save(ROOT/'muscle-control-classic-navy-structured.pptx')
# A prose companion makes text review possible without PowerPoint.
parts=['# Poster content\n\nAbstract source: ../docs/icra2027/main.tex (matches submission PDF).\n']
for name in ['ABSTRACT','INTRODUCTION','FIGURE 1 · PROJECT OVERVIEW','RELATED WORK','FORMULATION / MODEL','RESULTS','CONCLUSION','REFERENCES']:
    parts.append('\n## '+name+'\n')
# Include all visible editable text in reading/creation order.
parts=['# Structured poster text\n\nAbstract condensed from the ICRA 2027 submission.\n']
for s in slide.shapes:
    if s.has_text_frame and s.text.strip():parts.append(s.text+'\n')
(ROOT/'content/structured-poster.txt').write_text('\n'.join(parts))
print(ROOT/'muscle-control-classic-navy-structured.pptx')
